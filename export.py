"""
Designed Word and PDF exports for the BRIEF report and the guide audit.

Both formats are rendered from one intermediate document model built from the
result JSON, so they always agree with each other. The model is a flat list of
elements:

  ("title", heading, subtitle_lines)          cover block
  ("score", number, label, colour, caption)   big score with caption
  ("h", level, text)                          heading, level 1-3
  ("p", text)                                 paragraph; **bold** and _italic_ supported
  ("callout", tone, label, text)              boxed text; tone accent|danger|amber|ok|neutral
  ("kv", [(key, value), ...])                 two-column facts table
  ("table", headers, rows, widths, tone_col)  data table; tone_col colours a column by its text
  ("bullets", items) / ("numbered", items)    lists
  ("small", text)                             footnote-sized text
  ("pagebreak",)
"""

from __future__ import annotations

import io
import re
import time
from html import escape
from typing import Any, Iterator

ACCENT = "3A34B8"
DANGER = "C33C3C"
AMBER = "A86F12"
OK = "21855A"
INK = "1B1A17"
INK2 = "514E48"
MUTED = "8A867D"
SURFACE = "F3F1EA"
SURFACE_ALT = "F9F8F4"
HEADER_FILL = "ECE9F8"
BORDER = "D3CDBF"

TONE_COLOUR = {"accent": ACCENT, "danger": DANGER, "amber": AMBER, "ok": OK, "neutral": MUTED}

# Words that colour a table cell when tone_col is set. Longer keys first so "not supported" beats "supported".
CELL_TONES = [
    ("not supported", DANGER), ("partly supported", AMBER), ("confirmed only", AMBER), ("confirmed_only", AMBER),
    ("insufficient", AMBER), ("compromised", DANGER), ("critical", DANGER), ("untested", DANGER), ("adequate", AMBER),
    ("fragile", AMBER), ("supported", OK), ("against", DANGER), ("tested", OK), ("strong", OK), ("change", DANGER),
    ("adjust", AMBER), ("medium", AMBER), ("high", DANGER), ("keep", OK), ("low", OK), ("for", OK),
]


_UNPRINTABLE = re.compile(r"[\u2500-\u25ff\U0001F000-\U0001FFFF\ufffd]")


def _s(value: Any, limit: int = 4000) -> str:
    """Collapse whitespace, drop glyphs the PDF fonts cannot draw, and cut at a word boundary."""
    text = _UNPRINTABLE.sub("", str(value if value is not None else ""))
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= limit:
        return text
    cut = text[:limit].rsplit(" ", 1)[0]
    return cut + "\u2026"


def _score_colour(score: Any) -> str:
    try:
        score = int(score)
    except (TypeError, ValueError):
        return MUTED
    return OK if score >= 70 else AMBER if score >= 45 else DANGER


def _cell_tone(text: str) -> str | None:
    key = text.strip().lower()
    for word, colour in CELL_TONES:
        if key == word or key.startswith(word + " ") or key.startswith(word + ":"):
            return colour
    return None


def _verdict_tone(verdict: Any) -> str:
    v = _s(verdict).lower()
    if v.startswith("not"):
        return "danger"
    if v.startswith(("partly", "insufficient")):
        return "amber"
    if v.startswith("supported"):
        return "ok"
    return "neutral"


def _decision_tone(decision: Any) -> str:
    d = _s(decision).lower()
    return "ok" if d == "keep" else "amber" if d == "adjust" else "danger"


def _src(f: dict[str, Any]) -> str:
    year = _s(f.get("year"), 4)
    source = _s(f.get("source"), 140)
    if year and year in source:
        return source
    return source + (f" · {year}" if year else "")


def _plain(text: Any, limit: int) -> str:
    """Raw AI answers carry markdown; strip headings, bullets, emphasis and code fences for print."""
    t = str(text if text is not None else "")
    t = re.sub(r"```.*?```", " ", t, flags=re.S)
    t = re.sub(r"^\s{0,3}#{1,6}\s*", "", t, flags=re.M)
    t = re.sub(r"^\s*[-*•]\s+", "", t, flags=re.M)
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"(?<!\w)[*_](.+?)[*_](?!\w)", r"\1", t)
    t = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1", t)
    return _s(t, limit)


_JUNK_TITLE = re.compile(r"follow us|facebook|@\w+|download\?|cookie|sign in|log in", re.I)


def _good_title(title: Any) -> bool:
    t = _s(title, 200)
    return bool(t) and t.lower() != "source" and not _JUNK_TITLE.search(t)


def _avoid(text: Any) -> str:
    """The 'do not ask' line without a duplicated lead-in."""
    t = _s(text, 300)
    return re.sub(r"^(do not ask|don't ask|avoid( asking)?)[:\s]*", "", t, flags=re.I).strip()


# ---------------------------------------------------------------------------
# Document model: report
# ---------------------------------------------------------------------------

def build_report(data: dict[str, Any]) -> list[tuple]:
    parsed = data.get("parsed") or {}
    conf = data.get("confidence") or {}
    gaps = data.get("gaps") or {}
    cont = data.get("contamination") or {}
    drift = data.get("temporal_drift") or {}
    arch = data.get("archaeology") or {}
    meth = data.get("methodology") or {}
    comp = data.get("competitor_intel") or {}
    dl = data.get("deliverables") or {}
    health = data.get("run_health") or {}
    assurance = data.get("assurance") or {}
    internal = data.get("internal_recommendation") or {}
    human_review = data.get("human_review") or {}
    clusters = data.get("clusters") or {}
    el: list[tuple] = []

    mode = "UX research" if parsed.get("research_mode") == "ux" else "Market research"
    el.append(("title", "Research quality-assurance review", [
        _s(parsed.get("core_question") or parsed.get("research_objective"), 220),
        f"{mode}  ·  {time.strftime('%d %B %Y')}",
    ]))
    el.append(("score", conf.get("confidence_score", "n/a"), _s(conf.get("confidence_label")),
               _score_colour(conf.get("confidence_score")), "Heuristic research-design review indicator, out of 100"))
    if internal:
        tone = "danger" if internal.get("decision") == "hold" else "amber" if internal.get("decision") == "proceed_with_changes" else "accent"
        detail = _s(internal.get("recommended_action"), 500)
        reasons = "; ".join(_s(x, 300) for x in internal.get("reasons") or [])
        el.append(("callout", tone, f"Internal workflow recommendation: {_s(internal.get('label'))}",
                   detail + (f" Reasons: {reasons}" if reasons else "") + f" {_s(internal.get('notice'), 500)}"))
    if conf.get("headline"):
        el.append(("callout", "neutral", "", _s(conf["headline"])))
    if conf.get("score_rationale"):
        el.append(("p", _s(conf["score_rationale"])))
    if assurance:
        notice = _s(assurance.get("metric_notice"), 500)
        level = _s(assurance.get("assurance_level") or "limited").capitalize()
        el.append(("callout", "amber", f"{level} assurance · human review required", notice))
        limitations = [_s(x, 400) for x in assurance.get("limitations", []) if _s(x)]
        if limitations:
            el.append(("h", 3, "Limitations"))
            el.append(("bullets", limitations))
        decisions = [_s(x, 400) for x in assurance.get("required_reviewer_decisions", []) if _s(x)]
        if decisions:
            el.append(("h", 3, "Reviewer sign-off"))
            el.append(("numbered", decisions))

    if human_review:
        status = "Complete" if human_review.get("complete") else "Incomplete"
        el.append(("callout", "neutral", "Researcher adjudication", f"{status}. Reviewer: {_s(human_review.get('reviewer'))}. Reviewed: {_s(human_review.get('reviewed_at'))}."))
        review_rows = [[_s(x.get("finding_id")), _s(x.get("decision")), _s(x.get("rationale"), 500), _s(x.get("amendment"), 500)]
                       for x in human_review.get("decisions") or []]
        if review_rows:
            el.append(("table", ["Finding", "Decision", "Rationale", "Amendment"], review_rows,
                       [0.16, 0.14, 0.38, 0.32], None))

    if health.get("quick_mode") or health.get("step_errors"):
        note = "Quick mode: one AI model, no web evidence. Scores are indicative. " if health.get("quick_mode") else ""
        if health.get("step_errors"):
            note += "Steps that used a fallback: " + ", ".join(_s(e.get("step")) for e in health["step_errors"]) + "."
        el.append(("callout", "amber", "Run health", note.strip()))

    sct = gaps.get("sample_cannot_test") or []
    if sct:
        text = " ".join(f"**{_s(x.get('hypothesis'))}** {_s(x.get('why'))}" for x in sct)
        if conf.get("score_capped"):
            text += f" {_s(conf['score_capped'])}"
        el.append(("callout", "danger", "The stated sample cannot test these hypotheses", text))
    kf = conf.get("key_finding") or {}
    if kf.get("statement"):
        src = "; ".join(_s(x) for x in kf.get("sources", []) if x)
        el.append(("callout", "accent", "Key finding", _s(kf["statement"]) + (f" _Sources: {src}_" if src else "")))
    risks = [_s(r) for r in conf.get("top_three_risks", []) if _s(r)]
    if risks:
        el.append(("h", 2, "Fix before fieldwork"))
        el.append(("numbered", risks))

    el.append(("pagebreak",))
    el.append(("h", 1, "The brief as read"))
    facts = [("Category", parsed.get("category")), ("Specific topic", parsed.get("topic")),
             ("Audience", parsed.get("target_audience")),
             ("Geography", parsed.get("geography")), ("Objective", parsed.get("research_objective")),
             ("Stated method", parsed.get("methodology_hints")), ("Recruits", parsed.get("sample_definition")),
             ("Fieldwork", parsed.get("fieldwork_locations")), ("Product", parsed.get("product_or_service"))]
    el.append(("kv", [(k, _s(v, 300)) for k, v in facts if _s(v) and _s(v) != "Not specified"]))
    hyps = [_s(h) for h in parsed.get("client_hypotheses", []) if _s(h)]
    if hyps:
        el.append(("h", 3, "Client hypotheses"))
        el.append(("numbered", hyps))

    el.append(("h", 1, "Hypothesis contamination"))
    if cont.get("overall_contamination_level"):
        el.append(("p", f"**Overall: {_s(cont['overall_contamination_level'])}.** {_s(cont.get('overall_explanation'))}"))
    assessed = cont.get("hypotheses_assessed") or []
    if assessed:
        rows = [[_s(h.get("hypothesis"), 240),
                 str(h.get("contamination_score")) if h.get("contamination_score") is not None else _s(h.get("score_label"), 40),
                 _s(h.get("score_label")),
                 _s(h.get("responses_matching"), 200)] for h in assessed]
        el.append(("table", ["Hypothesis", "Score", "Band", "Measured across the AI answers"], rows, [0.40, 0.09, 0.12, 0.39], 2))
        el.append(("small", "Score: a classifier marks each AI answer as presenting the idea as the main cause (1), as one factor "
                            "among several (0.6), disputing it (-0.5) or omitting it (0); the total is divided by the number of "
                            "answers. Bands: 0-25 Low, 26-50 Medium, 51-75 High, 76-100 Critical."))
        for h in assessed:
            el.append(("h", 3, _s(h.get("hypothesis"), 200)))
            if h.get("explanation"):
                el.append(("p", _s(h["explanation"])))
            quotes = [_s(q, 300) for q in h.get("evidence_quotes", []) if _s(q)]
            if quotes:
                el.append(("bullets", [f"\u201c{q}\u201d" for q in quotes[:4]]))
            if h.get("recommendation"):
                el.append(("p", f"**Recommendation.** {_s(h['recommendation'])}"))

    el.append(("h", 1, "What AI assumes about this topic"))
    if clusters.get("dominant_perspective"):
        el.append(("p", _s(clusters["dominant_perspective"])))
    dom = clusters.get("dominant_assumptions") or []
    if dom:
        rows = [[_s(a.get("theme"), 120), _s(a.get("frequency"), 60), _s(a.get("description"), 300)] for a in dom]
        el.append(("table", ["Assumption", "Frequency", "What AI assumes"], rows, [0.28, 0.14, 0.58], None))
    if gaps.get("audience_mismatch"):
        el.append(("h", 3, "Where AI's picture and your audience differ"))
        el.append(("p", _s(gaps["audience_mismatch"])))
    over = gaps.get("overrepresented") or []
    under = gaps.get("underrepresented") or []
    if over or under:
        rows = [["Over-represented", _s(x.get("perspective", x), 120), _s(x.get("explanation", ""), 300)] for x in over]
        rows += [["Under-represented", _s(x.get("perspective", x), 120), _s(x.get("explanation", ""), 300)] for x in under]
        el.append(("table", ["", "Perspective", "Why it matters"], rows, [0.20, 0.28, 0.52], None))
    unknowns = [_s(u.get("question", u) if isinstance(u, dict) else u) for u in gaps.get("unknown_unknowns") or []]
    unknowns = [u for u in unknowns if u]
    if unknowns:
        el.append(("h", 3, "Unknown unknowns to explore"))
        el.append(("bullets", unknowns))

    hevid = [he for he in arch.get("hypothesis_evidence") or [] if he.get("grounded")]
    if hevid or arch.get("dominant_narrative_origin"):
        el.append(("pagebreak",))
        el.append(("h", 1, "Published evidence"))
    for he in hevid:
        el.append(("h", 2, _s(he.get("hypothesis"), 200)))
        el.append(("callout", _verdict_tone(he.get("verdict")), f"Verdict: {_s(he.get('verdict'))}", _s(he.get("summary"))))
        for c in he.get("countries") or []:
            rows = [["For", _s(f.get("finding"), 700), _src(f)] for f in c.get("for") or []]
            rows += [["Against", _s(f.get("finding"), 700), _src(f)] for f in c.get("against") or []]
            if rows:
                el.append(("h", 3, f"{_s(c.get('country'))}  ·  {_s(c.get('verdict'))}"))
                el.append(("table", ["", "Finding", "Source"], rows, [0.12, 0.58, 0.30], 0))
    if arch.get("dominant_narrative_origin"):
        el.append(("h", 2, "Where these assumptions come from"))
        el.append(("p", _s(arch["dominant_narrative_origin"])))
        if arch.get("implication_for_research"):
            el.append(("p", _s(arch["implication_for_research"])))
        srcs = [s for s in arch.get("sources") or [] if s.get("url") and _good_title(s.get("title"))]
        if srcs:
            el.append(("bullets", [f"{_s(s.get('title'), 90)} \u2014 {_s(s.get('url'), 160)}" for s in srcs[:6]]))

    if drift.get("overall_drift_risk"):
        el.append(("h", 1, "Temporal drift"))
        el.append(("p", f"**Risk: {_s(drift['overall_drift_risk'])}.** {_s(drift.get('drift_explanation'))}"))
        stale = drift.get("stale_assumptions") or []
        if stale:
            rows = [[_s(x.get("assumption"), 160), _s(x.get("why_stale"), 260), _s(x.get("research_implication"), 220)] for x in stale]
            el.append(("table", ["Assumption", "What has changed", "Probe in fieldwork"], rows, [0.28, 0.40, 0.32], None))

    brands = comp.get("brands_mentioned") or []
    if brands:
        el.append(("h", 1, "Brands AI puts in the room"))
        rows = [[_s(b.get("brand"), 80), _s(b.get("frequency"), 60), _s(b.get("framing"), 160), _s(b.get("priming_risk"), 30)] for b in brands]
        el.append(("table", ["Brand", "Frequency", "How AI frames it", "Priming risk"], rows, [0.22, 0.14, 0.46, 0.18], 3))
        if comp.get("discussion_guide_implication"):
            el.append(("p", _s(comp["discussion_guide_implication"])))

    el.append(("pagebreak",))
    el.append(("h", 1, "Methodology"))
    fit = meth.get("method_fit") or {}
    if fit.get("decision"):
        el.append(("callout", _decision_tone(fit["decision"]),
                   f"Fit with the method in the brief ({_s(fit.get('stated_method'), 80)}): {_s(fit['decision'])}",
                   f"{_s(fit.get('reason'))} {_s(fit.get('scope_and_cost_note'))}".strip()))
    if meth.get("recommended_approach"):
        el.append(("p", f"**Recommended approach.** {_s(meth['recommended_approach'])}"))
    if meth.get("sample_design_notes"):
        el.append(("p", f"**Sample.** {_s(meth['sample_design_notes'])}"))
    breakdown = meth.get("methodology_breakdown") or []
    if breakdown:
        rows = [[_s(x.get("dimension"), 140), _s(x.get("recommended_method"), 120), _s(x.get("rationale"), 240)] for x in breakdown]
        el.append(("table", ["Dimension", "Method", "Why"], rows, [0.30, 0.25, 0.45], None))
    tests = meth.get("hypothesis_tests") or []
    if tests:
        el.append(("h", 3, "How to test each client hypothesis"))
        rows = [[_s(t.get("hypothesis"), 160), _s(t.get("how_to_test_it"), 280), _s(t.get("what_would_refute_it"), 220)] for t in tests]
        el.append(("table", ["Hypothesis", "How to test it", "Drop it if"], rows, [0.28, 0.40, 0.32], None))
    if meth.get("stimulus_material_warning"):
        el.append(("small", f"Stimulus: {_s(meth['stimulus_material_warning'])}"))

    note = dl.get("challenge_note") or {}
    if note.get("what_we_found"):
        el.append(("pagebreak",))
        el.append(("h", 1, "Challenge note for the client"))
        if note.get("title"):
            el.append(("h", 3, _s(note["title"], 160)))
        for key in ("opening", "what_we_found", "what_we_recommend", "closing"):
            if note.get(key):
                el.append(("p", _s(note[key])))
    probes = dl.get("discussion_guide_probes") or []
    if probes:
        el.append(("h", 1, "Discussion guide probes"))
        for g in probes:
            el.append(("h", 3, f"Tests: {_s(g.get('hypothesis'), 160)}"))
            if g.get("warm_up"):
                el.append(("p", f"**Warm-up.** {_s(g['warm_up'])}"))
            el.append(("bullets", [_s(p) for p in g.get("probes") or [] if _s(p)]))
            if g.get("avoid"):
                el.append(("small", f"Do not ask: {_avoid(g['avoid'])}"))
    tasks = [t for t in dl.get("task_scenarios") or [] if t.get("task")]
    if tasks:
        el.append(("h", 1, "Usability task scenarios"))
        rows = [[_s(t.get("task"), 240), _s(t.get("success_measure"), 200), _s(t.get("hypothesis_tested"), 140)] for t in tasks]
        el.append(("table", ["Task", "Success looks like", "Tests"], rows, [0.42, 0.34, 0.24], None))
    screener = dl.get("screener_criteria") or []
    if screener:
        el.append(("h", 1, "Screener criteria"))
        rows = [[_s(c.get("criterion"), 200), _s(c.get("quota_note"), 120), _s(c.get("why"), 240)] for c in screener]
        el.append(("table", ["Criterion", "Quota", "Why"], rows, [0.36, 0.22, 0.42], None))

    q = data.get("query_data") or {}
    answers = (q.get("base_responses") or []) + (q.get("persona_responses") or [])
    if answers:
        el.append(("pagebreak",))
        el.append(("h", 1, "Appendix: what AI actually said"))
        models = ", ".join(health.get("probe_models") or q.get("models") or [])
        el.append(("small", f"{len(answers)} answers from {models}. Each is cut to 500 characters; the full text is in the run log."))
        for i, r in enumerate(answers, 1):
            tag = f"[{_s(r.get('model'))}]" + (f" [{_s(r.get('persona'))}]" if r.get("persona") else "")
            el.append(("h", 3, f"{i}. {tag} {_s(r.get('prompt'), 200)}"))
            el.append(("small", _plain(r.get("response"), 500)))
    return el


# ---------------------------------------------------------------------------
# Document model: guide audit
# ---------------------------------------------------------------------------

def build_audit(data: dict[str, Any]) -> list[tuple]:
    sm = data.get("summary") or {}
    assurance = data.get("assurance") or {}
    human_review = data.get("human_review") or {}
    internal = data.get("internal_recommendation") or {}
    items = data.get("items") or []
    cov = data.get("coverage") or {}
    el: list[tuple] = []
    el.append(("title", "Guide and questionnaire audit", [
        f"{_s(data.get('instrument_type', 'instrument')).capitalize()}  ·  {sm.get('items_total', 0)} items  ·  {time.strftime('%d %B %Y')}",
    ]))
    el.append(("score", sm.get("score", "n/a"), _s(sm.get("label")), _score_colour(sm.get("score")), "Heuristic wording-review indicator, out of 100"))
    if internal:
        tone = "danger" if internal.get("decision") == "hold" else "amber" if internal.get("decision") == "proceed_with_changes" else "accent"
        detail = _s(internal.get("recommended_action"), 500)
        reasons = "; ".join(_s(x, 300) for x in internal.get("reasons") or [])
        el.append(("callout", tone, f"Internal workflow recommendation: {_s(internal.get('label'))}",
                   detail + (f" Reasons: {reasons}" if reasons else "") + f" {_s(internal.get('notice'), 500)}"))
    el.append(("p", f"{sm.get('items_flagged', 0)} of {sm.get('items_total', 0)} items flagged: {sm.get('high', 0)} high, "
                    f"{sm.get('medium', 0)} medium, {sm.get('low', 0)} low. {sm.get('items_confirming', 0)} restate a client hypothesis."))
    el.append(("small", "Score starts at 100 and loses 12 per high, 5 per medium and 1 per low issue, scaled for short "
                        "instruments. It measures wording, not study design."))
    if human_review:
        status = "Complete" if human_review.get("complete") else "Incomplete"
        el.append(("callout", "neutral", "Researcher adjudication", f"{status}. Reviewer: {_s(human_review.get('reviewer'))}. Reviewed: {_s(human_review.get('reviewed_at'))}."))
        review_rows = [[_s(x.get("finding_id")), _s(x.get("decision")), _s(x.get("rationale"), 500), _s(x.get("amendment"), 500)]
                       for x in human_review.get("decisions") or []]
        if review_rows:
            el.append(("table", ["Finding", "Decision", "Rationale", "Amendment"], review_rows,
                       [0.16, 0.14, 0.38, 0.32], None))

    worst = [i for i in items if any(x.get("severity") == "high" for x in i.get("issues") or [])][:3]
    if worst:
        el.append(("callout", "danger", "Fix these first", " ".join(
            f"**{_s(i.get('id'))}** {_s(i.get('text'), 160)} \u2014 {_s(i['issues'][0].get('explanation'), 160)}." for i in worst)))
    if sm.get("untested_hypotheses"):
        el.append(("callout", "danger", "Client hypotheses this instrument never tests", "; ".join(_s(h) for h in sm["untested_hypotheses"])))
    if sm.get("confirmed_only_hypotheses"):
        el.append(("callout", "amber", "Hypotheses it can only confirm", "; ".join(_s(h) for h in sm["confirmed_only_hypotheses"])))

    el.append(("h", 1, "Question by question"))
    rows = []
    for i in items:
        issues = i.get("issues") or []
        severities = [_s(x.get("severity")).lower() for x in issues]
        sev = "high" if "high" in severities else "medium" if "medium" in severities else "low" if issues else "clean"
        problems = "; ".join(f"{_s(x.get('type')).replace('_', ' ')}: {_s(x.get('explanation'), 140)}" for x in issues) or "No issues"
        rows.append([_s(i.get("id"), 8), _s(i.get("text"), 260), sev, problems, _s(i.get("rewrite"), 260) or "\u2014"])
    el.append(("table", ["#", "As written", "Severity", "Problems", "Neutral version"], rows, [0.06, 0.28, 0.10, 0.28, 0.28], 2))

    hyps = cov.get("hypotheses") or []
    if hyps:
        el.append(("h", 1, "Hypothesis coverage"))
        rows = [[_s(h.get("hypothesis"), 200), _s(h.get("coverage")).replace("_", " "),
                 ", ".join(_s(x) for x in h.get("item_ids") or []) or "none", _s(h.get("note"), 220)] for h in hyps]
        el.append(("table", ["Hypothesis", "Coverage", "Items", "Note"], rows, [0.34, 0.16, 0.14, 0.36], 1))
    missing = cov.get("missing_questions") or []
    if missing:
        el.append(("h", 3, "Questions to add"))
        el.append(("bullets", [f"{_s(q.get('question'))} _(serves: {_s(q.get('serves'), 120)}; place after {_s(q.get('place_after'), 20)})_"
                               for q in missing]))

    el.append(("pagebreak",))
    el.append(("h", 1, "Clean copy"))
    el.append(("small", "The instrument with every rewrite applied. Restore your own voice where a rewrite went flat."))
    section = None
    for i in items:
        if i.get("section") and i["section"] != section:
            section = i["section"]
            el.append(("h", 3, _s(section, 120)))
        el.append(("p", f"**{_s(i.get('id'))}.** {_s(i.get('rewrite') or i.get('text'))}"))
    return el


# ---------------------------------------------------------------------------
# Inline formatting
# ---------------------------------------------------------------------------

def _runs(text: str) -> Iterator[tuple[str, bool, bool]]:
    """Yield (chunk, bold, italic) for **bold** and _italic_ markers."""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|(?<![A-Za-z0-9])_(.+?)_(?![A-Za-z0-9])", text):
        if m.start() > pos:
            yield text[pos:m.start()], False, False
        if m.group(1) is not None:
            yield m.group(1), True, False
        else:
            yield m.group(2), False, True
        pos = m.end()
    if pos < len(text):
        yield text[pos:], False, False


def _pdf_markup(text: str) -> str:
    out = []
    for chunk, bold, italic in _runs(text):
        piece = escape(chunk)
        if bold:
            piece = f"<b>{piece}</b>"
        if italic:
            piece = f"<i>{piece}</i>"
        out.append(piece)
    return "".join(out)


# ---------------------------------------------------------------------------
# PDF renderer
# ---------------------------------------------------------------------------

def render_pdf(elements: list[tuple], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    def hexc(h: str):
        return colors.HexColor("#" + h)

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=9.5, leading=13.5, textColor=hexc(INK), spaceAfter=6, alignment=TA_LEFT)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=11, textColor=hexc(INK2), spaceAfter=4)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11.5, spaceAfter=0)
    cell_head = ParagraphStyle("cellhead", parent=cell, fontName="Helvetica-Bold", textColor=hexc(ACCENT), fontSize=7.5)
    headings = {
        1: ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=hexc(INK), spaceBefore=14, spaceAfter=8),
        2: ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=13, leading=17, textColor=hexc(ACCENT), spaceBefore=12, spaceAfter=5),
        3: ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=10.5, leading=14, textColor=hexc(INK), spaceBefore=9, spaceAfter=3),
    }
    title_style = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=26, leading=30, textColor=hexc(INK), spaceAfter=6)
    subtitle = ParagraphStyle("subtitle", parent=body, fontSize=10.5, leading=15, textColor=hexc(INK2))
    eyebrow = ParagraphStyle("eyebrow", parent=small, fontName="Helvetica-Bold", textColor=hexc(ACCENT), fontSize=8)
    width = A4[0] - 36 * mm
    pad = [("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
           ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]

    story: list = []
    for e in elements:
        kind = e[0]
        if kind == "title":
            story.append(Paragraph("BRIEF  ·  BIAS &amp; RESEARCH INTELLIGENCE EVALUATION FRAMEWORK", eyebrow))
            story.append(Paragraph(_pdf_markup(e[1]), title_style))
            for line in e[2]:
                if line:
                    story.append(Paragraph(_pdf_markup(line), subtitle))
            story.append(Spacer(1, 10))
        elif kind == "score":
            _, number, label, colour, caption = e
            num = Paragraph(f'<font size="34" color="#{colour}"><b>{escape(str(number))}</b></font>', body)
            lab = Paragraph(f'<font size="12" color="#{colour}"><b>{escape(label)}</b></font><br/>'
                            f'<font size="8" color="#{MUTED}">{escape(caption)}</font>', body)
            t = Table([[num, lab]], colWidths=[width * 0.22, width * 0.78])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BACKGROUND", (0, 0), (-1, -1), hexc(SURFACE)),
                                   ("BOX", (0, 0), (-1, -1), 0.5, hexc(BORDER)), ("LEFTPADDING", (0, 0), (-1, -1), 10),
                                   ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
            story.append(t)
            story.append(Spacer(1, 10))
        elif kind == "h":
            story.append(Paragraph(_pdf_markup(e[2]), headings[min(e[1], 3)]))
        elif kind == "p":
            story.append(Paragraph(_pdf_markup(e[1]), body))
        elif kind == "small":
            story.append(Paragraph(_pdf_markup(e[1]), small))
        elif kind == "callout":
            _, tone, label, text = e
            colour = TONE_COLOUR.get(tone, MUTED)
            inner = (f'<font size="7.5" color="#{colour}"><b>{escape(label).upper()}</b></font><br/>' if label else "") + _pdf_markup(text)
            t = Table([[Paragraph(inner, body)]], colWidths=[width])
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), hexc(SURFACE)), ("LINEBEFORE", (0, 0), (0, -1), 2.5, hexc(colour)),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                                   ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story.append(KeepTogether([t, Spacer(1, 8)]))
        elif kind == "kv":
            rows = [[Paragraph(f"<b>{escape(k)}</b>", cell), Paragraph(_pdf_markup(v), cell)] for k, v in e[1]]
            if rows:
                t = Table(rows, colWidths=[width * 0.22, width * 0.78])
                t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.4, hexc(BORDER)), ("VALIGN", (0, 0), (-1, -1), "TOP")] + pad))
                story.append(t)
                story.append(Spacer(1, 8))
        elif kind == "table":
            _, headers, rows, widths, tone_col = e
            table_data = [[Paragraph(escape(x).upper(), cell_head) for x in headers]]
            for row in rows:
                cells = []
                for ci, value in enumerate(row):
                    tone = _cell_tone(value) if tone_col is not None and ci == tone_col else None
                    txt = _pdf_markup(value)
                    if tone:
                        txt = f'<font color="#{tone}"><b>{txt}</b></font>'
                    cells.append(Paragraph(txt, cell))
                table_data.append(cells)
            t = Table(table_data, colWidths=[width * w for w in widths], repeatRows=1)
            style = [("LINEBELOW", (0, 0), (-1, 0), 0.8, hexc(ACCENT)), ("LINEBELOW", (0, 1), (-1, -1), 0.3, hexc(BORDER)),
                     ("VALIGN", (0, 0), (-1, -1), "TOP")] + pad
            for ri in range(2, len(table_data), 2):
                style.append(("BACKGROUND", (0, ri), (-1, ri), hexc(SURFACE_ALT)))
            t.setStyle(TableStyle(style))
            story.append(t)
            story.append(Spacer(1, 8))
        elif kind in ("bullets", "numbered"):
            items = [ListItem(Paragraph(_pdf_markup(x), cell), leftIndent=14) for x in e[1] if x]
            if items:
                if kind == "numbered":
                    story.append(ListFlowable(items, bulletType="1", leftIndent=14, bulletFontName="Helvetica-Bold",
                                              bulletFontSize=9, bulletColor=hexc(ACCENT)))
                else:
                    story.append(ListFlowable(items, bulletType="bullet", start="\u2022", leftIndent=14,
                                              bulletFontSize=8, bulletColor=hexc(ACCENT)))
                story.append(Spacer(1, 6))
        elif kind == "pagebreak":
            story.append(PageBreak())

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(hexc(BORDER))
        canvas.line(18 * mm, 16 * mm, A4[0] - 18 * mm, 16 * mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(hexc(MUTED))
        canvas.drawString(18 * mm, 11 * mm, f"{title}  ·  Generated by BRIEF")
        canvas.drawRightString(A4[0] - 18 * mm, 11 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                            bottomMargin=22 * mm, title=title, author="BRIEF")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# DOCX renderer
# ---------------------------------------------------------------------------

def render_docx(elements: list[tuple], title: str) -> bytes:
    import docx
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def rgb(h: str) -> RGBColor:
        return RGBColor.from_string(h)

    def shade(table_cell, fill: str) -> None:
        pr = table_cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        pr.append(shd)

    def borders(table, colour: str, size: int = 4, inside: bool = True) -> None:
        tbl_pr = table._tbl.tblPr
        el = OxmlElement("w:tblBorders")
        edges = ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ())
        for edge in edges:
            b = OxmlElement(f"w:{edge}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), str(size))
            b.set(qn("w:color"), colour)
            el.append(b)
        tbl_pr.append(el)

    def left_bar(table_cell, colour: str) -> None:
        pr = table_cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "24")
        left.set(qn("w:color"), colour)
        b.append(left)
        pr.append(b)

    def add_runs(paragraph, text: str, size: float | None = None, colour: str | None = None, bold_all: bool = False) -> None:
        for chunk, bold, italic in _runs(text):
            run = paragraph.add_run(chunk)
            run.bold = bold or bold_all
            run.italic = italic
            if size:
                run.font.size = Pt(size)
            if colour:
                run.font.color.rgb = rgb(colour)

    def spacer() -> None:
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    def fix_widths(table, widths_emu: list[int]) -> None:
        """Word and LibreOffice only honour widths when the layout is fixed and every cell is set."""
        table.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
        for ci, w in enumerate(widths_emu):
            table.columns[ci].width = w
            for row in table.rows:
                row.cells[ci].width = w

    def numbered_paragraph(index: int, text: str) -> None:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.space_after = Pt(3)
        add_runs(p, f"{index}.\t", size=9.5, colour=ACCENT, bold_all=True)
        add_runs(p, text, size=9.5)

    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb(INK)
    for level, size, colour in ((1, 17, INK), (2, 13, ACCENT), (3, 11, INK)):
        st = document.styles[f"Heading {level}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(colour)
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(2)
        add_runs(section.footer.paragraphs[0], f"{title}  ·  Generated by BRIEF", size=8, colour=MUTED)
    usable = Cm(17)

    for e in elements:
        kind = e[0]
        if kind == "title":
            add_runs(document.add_paragraph(), "BRIEF  ·  BIAS & RESEARCH INTELLIGENCE EVALUATION FRAMEWORK", size=8, colour=ACCENT, bold_all=True)
            add_runs(document.add_paragraph(), e[1], size=24, bold_all=True)
            for line in e[2]:
                if line:
                    add_runs(document.add_paragraph(), line, size=10.5, colour=INK2)
        elif kind == "score":
            _, number, label, colour, caption = e
            t = document.add_table(rows=1, cols=2)
            borders(t, BORDER, 4, inside=False)
            for c in t.rows[0].cells:
                shade(c, SURFACE)
            fix_widths(t, [Cm(3.5), usable - Cm(3.5)])
            add_runs(t.cell(0, 0).paragraphs[0], str(number), size=30, colour=colour, bold_all=True)
            add_runs(t.cell(0, 1).paragraphs[0], label, size=12, colour=colour, bold_all=True)
            add_runs(t.cell(0, 1).add_paragraph(), caption, size=8, colour=MUTED)
            spacer()
        elif kind == "h":
            level = min(e[1], 3)
            size, colour = {1: (17, INK), 2: (13, ACCENT), 3: (11, INK)}[level]
            add_runs(document.add_heading(level=level), e[2], size=size, colour=colour, bold_all=True)
        elif kind == "p":
            add_runs(document.add_paragraph(), e[1])
        elif kind == "small":
            add_runs(document.add_paragraph(), e[1], size=8.5, colour=INK2)
        elif kind == "callout":
            _, tone, label, text = e
            colour = TONE_COLOUR.get(tone, MUTED)
            t = document.add_table(rows=1, cols=1)
            c = t.cell(0, 0)
            shade(c, SURFACE)
            borders(t, BORDER, 4, inside=False)
            left_bar(c, colour)
            p = c.paragraphs[0]
            if label:
                add_runs(p, label.upper(), size=7.5, colour=colour, bold_all=True)
                p = c.add_paragraph()
            add_runs(p, text)
            spacer()
        elif kind == "kv":
            if not e[1]:
                continue
            t = document.add_table(rows=0, cols=2)
            borders(t, BORDER, 2)
            for k, v in e[1]:
                row = t.add_row()
                add_runs(row.cells[0].paragraphs[0], k, size=9, bold_all=True)
                add_runs(row.cells[1].paragraphs[0], v, size=9.5)
            fix_widths(t, [Cm(3.8), usable - Cm(3.8)])
            spacer()
        elif kind == "table":
            _, headers, rows, widths, tone_col = e
            t = document.add_table(rows=1, cols=len(headers))
            borders(t, BORDER, 2)
            for ci, header in enumerate(headers):
                c = t.rows[0].cells[ci]
                shade(c, HEADER_FILL)
                add_runs(c.paragraphs[0], header.upper(), size=7.5, colour=ACCENT, bold_all=True)
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for ci, value in enumerate(row):
                    tone = _cell_tone(value) if tone_col is not None and ci == tone_col else None
                    add_runs(cells[ci].paragraphs[0], value, size=8.5, colour=tone, bold_all=bool(tone))
                    if ri % 2 == 1:
                        shade(cells[ci], SURFACE_ALT)
            fix_widths(t, [int(usable * w) for w in widths])
            spacer()
        elif kind == "numbered":
            for index, item in enumerate([x for x in e[1] if x], 1):
                numbered_paragraph(index, item)
        elif kind == "bullets":
            for item in e[1]:
                if item:
                    add_runs(document.add_paragraph(style="List Bullet"), item, size=9.5)
        elif kind == "pagebreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    document.core_properties.title = title
    document.core_properties.author = "BRIEF"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def export_document(kind: str, data: dict[str, Any], fmt: str) -> tuple[bytes, str, str]:
    """Return (payload, mimetype, extension) for kind in {report, audit} and fmt in {docx, pdf}."""
    if kind == "audit":
        elements, title = build_audit(data), "BRIEF guide and questionnaire audit"
    else:
        elements, title = build_report(data), "BRIEF research quality-assurance review"
    if fmt == "pdf":
        return render_pdf(elements, title), "application/pdf", "pdf"
    if fmt == "docx":
        return (render_docx(elements, title),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx")
    raise ValueError("Unknown format. Use docx or pdf.")
