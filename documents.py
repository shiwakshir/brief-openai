"""Resource-bounded plain-text extraction from uploaded PDF, DOCX, TXT and MD files."""

from __future__ import annotations

import io
import zipfile

import config


class UnsupportedFileType(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        if not data.startswith(b"%PDF-"):
            raise UnsupportedFileType("The file does not appear to be a PDF.")
        text = _from_pdf(data)
    elif name.endswith(".docx"):
        if not data.startswith(b"PK"):
            raise UnsupportedFileType("The file does not appear to be a DOCX document.")
        text = _from_docx(data)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileType("Unsupported file type. Upload a PDF, DOCX, TXT or MD file.")
    return _tidy(text)


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=True)
    if len(reader.pages) > config.MAX_PDF_PAGES:
        raise ValueError(f"PDF exceeds the {config.MAX_PDF_PAGES}-page limit.")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data: bytes) -> str:
    import docx

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > config.MAX_DOCX_EXPANDED_BYTES:
                raise ValueError("DOCX expands beyond the configured safety limit.")
            if any(item.file_size > 100 * max(item.compress_size, 1) for item in archive.infolist()):
                raise ValueError("DOCX contains a suspiciously compressed component.")
    except zipfile.BadZipFile as exc:
        raise UnsupportedFileType("The DOCX file is invalid.") from exc

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _tidy(text: str) -> str:
    text = "\n".join(line.rstrip() for line in text.splitlines())
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    return text.strip()
